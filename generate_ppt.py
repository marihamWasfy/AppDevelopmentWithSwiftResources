from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# Create a presentation object
prs = Presentation()

# Slide layouts
title_slide_layout = prs.slide_layouts[0]  # Title Slide
content_slide_layout = prs.slide_layouts[1]  # Title and Content

# Function to add a slide with title, content, and notes
def add_slide(title, content_bullets, background_color=None, notes_text=""):
    slide = prs.slides.add_slide(content_slide_layout)
    shapes = slide.shapes
    
    # Set background color if provided
    if background_color:
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = background_color
    
    # Title
    title_shape = shapes.title
    title_shape.text = title
    title_shape.text_frame.paragraphs[0].font.size = Pt(40)
    title_shape.text_frame.paragraphs[0].font.bold = True
    title_shape.text_frame.paragraphs[0].font.color.rgb = RGBColor(178, 34, 34)  # Firebrick red
    
    # Content
    body_shape = shapes.placeholders[1]
    text_frame = body_shape.text_frame
    text_frame.clear()
    
    for bullet in content_bullets:
        p = text_frame.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(20)
        p.font.color.rgb = RGBColor(0, 100, 0)  # Dark green
    
    # Add notes
    if notes_text:
        notes_slide = slide.notes_slide
        text_frame = notes_slide.notes_text_frame
        text_frame.text = notes_text

# Slide 1: Title Slide
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "🎄 Magical Christmas Party 🎄"
subtitle.text = "A Festive Celebration at Our Church\nWhere Every Child Gets a Passport to Fun!"
title.text_frame.paragraphs[0].font.size = Pt(54)
title.text_frame.paragraphs[0].font.bold = True
title.text_frame.paragraphs[0].font.color.rgb = RGBColor(220, 20, 60)  # Crimson
subtitle.text_frame.paragraphs[0].font.size = Pt(24)
subtitle.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 128, 0)  # Green

# Slide 2: Welcome & Introduction
add_slide(
    "Welcome to Our Christmas Adventure! 🎅",
    [
        "Get ready for a magical journey!",
        "Fun activities, crafts, and surprises await",
        "Every child will receive a special passport",
        "Let's make this Christmas unforgettable!"
    ],
    RGBColor(255, 250, 240),  # Floral white background
    "Animation: Entrance effect with snowflakes. Image ideas: Santa waving, snowflakes, Christmas tree."
)

# Slide 3: Step 1 - Registration for Kids
add_slide(
    "Step 1: Registration - Your Passport to Fun! 🎫",
    [
        "Register your kids while purchasing tickets",
        "We'll prepare custom 'passports' for each child",
        "Passports = Official ID cards for the event",
        "Make sure to provide accurate names and ages!"
    ],
    RGBColor(240, 255, 240),  # Honeydew background
    "Animation: Slide in from left. Image ideas: Passport with Christmas stamp, kids registering at a desk."
)

# Slide 4: Step 2 - WhatsApp Group for Parents
add_slide(
    "Step 2: Join Our Christmas Community! 📱",
    [
        "Parents: Join our WhatsApp group!",
        "Get event updates and reminders",
        "Share gift ideas and coordinate with other families",
        "Invite your friends to join the celebration!"
    ],
    RGBColor(255, 250, 205),  # Lemon chiffon background
    "Animation: Zoom in effect. Image ideas: WhatsApp logo with Christmas theme, parents chatting on phones."
)

# Slide 5: Step 3 - Passport Distribution
add_slide(
    "Step 3: Passport Pick-Up Day! 🎉",
    [
        "Arrive at the party venue",
        "Kids receive their official passports at check-in",
        "Passports unlock all the fun activities",
        "Don't lose it - you'll need it at every station!"
    ],
    RGBColor(255, 240, 245),  # Lavender blush background
    "Animation: Appear with sparkle effect. Image ideas: Kids holding colorful passports, check-in desk with Santa helper."
)

# Slide 6: Step 4 - Slot Division & Story Time
add_slide(
    "Step 4: Story Time with Saint Nicolas! 📖✨",
    [
        "Kids divided into groups of 10 for personalized fun",
        "Each group hears the magical Saint Nicolas story",
        "Learn about the spirit of giving and kindness",
        "Get ready for the adventure ahead!"
    ],
    RGBColor(240, 248, 255),  # Alice blue background
    "Animation: Fade in with twinkling stars. Image ideas: Saint Nicolas telling story to kids sitting in a circle, storybook with illustrations."
)

# Slide 7: Step 5 - Art & Crafts Stations
add_slide(
    "Step 5: Crafty Elves Workshop! 🎨🖌️",
    [
        "4 exciting craft stations await!",
        "Each station: 10-15 minutes of creativity",
        "Make ornaments, cards, decorations, and more",
        "Volunteers guide kids through each activity"
    ],
    RGBColor(255, 245, 238),  # Seashell background
    "Animation: Wipe transition with craft tools. Image ideas: Kids painting, gluing, cutting paper, colorful craft supplies."
)

# Slide 8: Volunteer Tips for Craft Stations
add_slide(
    "Volunteer Tips: Craft Station Success! 🧑‍🎨",
    [
        "Keep activities simple and age-appropriate",
        "Have extra supplies ready for enthusiastic crafters",
        "Encourage creativity - there's no 'wrong' way!",
        "Help kids stamp their passports at each station"
    ],
    RGBColor(250, 250, 210),  # Light goldenrod yellow background
    "Animation: Checklist appear one by one. Image ideas: Volunteer helping child, craft station setup, passport being stamped."
)

# Slide 9: Step 6 - Tote Bag & Film Time
add_slide(
    "Step 6: Movie Magic & Tote Bags! 🎬🎒",
    [
        "Each child receives a special tote bag",
        "Store all your amazing crafts inside!",
        "Watch the finale of 'Christmas Factory' film",
        "Relax, enjoy snacks, and celebrate your creations"
    ],
    RGBColor(255, 248, 220),  # Cornsilk background
    "Animation: Curtain open effect. Image ideas: Kids with tote bags full of crafts, movie screen showing Christmas film, popcorn."
)

# Slide 10: Step 7 - Gift Wrapping Station
add_slide(
    "Step 7: Wrap It Up with Love! 🎁💝",
    [
        "Kids wrap the gifts they brought to share",
        "Learn the joy of giving to others",
        "Gifts stored safely in the 'Kasha Room'",
        "Beautiful wrapping paper and ribbons provided!"
    ],
    RGBColor(255, 240, 245),  # Lavender blush background
    "Animation: Gift boxes appear with bows. Image ideas: Kids wrapping presents, colorful wrapping paper, gift tags, 'Kasha Room' sign."
)

# Slide 11: Event Highlights Recap
add_slide(
    "What an Amazing Day! 🌟",
    [
        "✓ Passports collected and stamped",
        "✓ Saint Nicolas story shared",
        "✓ 4 craft stations completed",
        "✓ Tote bags filled with treasures",
        "✓ Gifts wrapped with love"
    ],
    RGBColor(240, 255, 255),  # Azure background
    "Animation: Checkmarks appear with ding sound. Image ideas: Collage of all activities, happy kids, Christmas decorations."
)

# Slide 12: Thank You & Future Events
slide = prs.slides.add_slide(content_slide_layout)
shapes = slide.shapes
background = slide.background
fill = background.fill
fill.solid()
fill.fore_color.rgb = RGBColor(255, 250, 240)  # Floral white

title_shape = shapes.title
title_shape.text = "Thank You for Celebrating with Us! 🙏🎄"
title_shape.text_frame.paragraphs[0].font.size = Pt(36)
title_shape.text_frame.paragraphs[0].font.bold = True
title_shape.text_frame.paragraphs[0].font.color.rgb = RGBColor(178, 34, 34)

body_shape = shapes.placeholders[1]
text_frame = body_shape.text_frame
text_frame.clear()

content = [
    "We hope you had a magical time!",
    "Share your photos on our WhatsApp group",
    "Stay tuned for more exciting church events",
    "Merry Christmas and God bless! 🎅✨"
]

for bullet in content:
    p = text_frame.add_paragraph()
    p.text = bullet
    p.level = 0
    p.font.size = Pt(22)
    p.font.color.rgb = RGBColor(0, 100, 0)

notes_slide = slide.notes_slide
text_frame = notes_slide.notes_text_frame
text_frame.text = "Animation: Fade out with sparkles. Image ideas: Group photo of all kids, 'Merry Christmas' banner, church logo."

# Save the presentation
prs.save('christmas_party_presentation.pptx')
print("✅ PowerPoint presentation created successfully!")
print("📁 File saved as: christmas_party_presentation.pptx")
print("🎄 Total slides: 12")
print("\n🎨 Design features:")
print("   - Festive red, green, and gold color scheme")
print("   - Large, readable fonts for all ages")
print("   - Bullet points for easy scanning")
print("   - Notes with animation and image suggestions")
print("\n💡 Next steps:")
print("   1. Open the file in PowerPoint or Google Slides")
print("   2. Add images/clipart as suggested in slide notes")
print("   3. Apply animations as described in notes")
print("   4. Customize colors/fonts to match your church branding")
print("   5. Add background music or embedded videos if desired")
